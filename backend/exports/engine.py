import csv
import json
import os
from pathlib import Path

from django.conf import settings

from datahub.models import ExtractedRecord
from .models import ExportJob

HEADERS = [
    "id", "entity_type", "location_id", "source_domain", "source_url",
    "status", "quality_score", "confidence", "canonical_key", "fingerprint",
    "payload", "normalized_payload", "evidence", "created_at",
]

MIME_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    "jsonl": "application/x-ndjson",
    "txt": "text/plain",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}

EXTENSIONS = {fmt: fmt for fmt in MIME_TYPES}


def _queryset(filters):
    qs = ExtractedRecord.objects.select_related("entity_type").order_by("pk")
    filters = filters or {}
    if filters.get("entity_type_id"):
        qs = qs.filter(entity_type_id=filters["entity_type_id"])
    if filters.get("location_id"):
        qs = qs.filter(location_id=filters["location_id"])
    if filters.get("source_domain"):
        qs = qs.filter(source_domain=filters["source_domain"])
    if filters.get("status"):
        qs = qs.filter(status=filters["status"])
    if filters.get("quality_min") not in (None, ""):
        qs = qs.filter(quality_score__gte=filters["quality_min"])
    if filters.get("created_after"):
        qs = qs.filter(created_at__gte=filters["created_after"])
    if filters.get("created_before"):
        qs = qs.filter(created_at__lte=filters["created_before"])
    return qs


def _row(record):
    return [
        record.pk,
        record.entity_type.slug if record.entity_type_id else "",
        record.location_id or "",
        record.source_domain,
        record.source_url,
        record.status,
        str(record.quality_score),
        str(record.confidence),
        record.canonical_key,
        record.fingerprint,
        json.dumps(record.payload, ensure_ascii=False, sort_keys=True),
        json.dumps(record.normalized_payload, ensure_ascii=False, sort_keys=True),
        json.dumps(record.evidence, ensure_ascii=False, sort_keys=True),
        record.created_at.isoformat(),
    ]


def _iter_rows(filters):
    for record in _queryset(filters).iterator(chunk_size=500):
        yield _row(record)


def _write_csv(path, filters):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for row in _iter_rows(filters):
            writer.writerow(row)


def _write_jsonl(path, filters):
    with path.open("w", encoding="utf-8") as handle:
        for row in _iter_rows(filters):
            handle.write(json.dumps(dict(zip(HEADERS, row)), ensure_ascii=False) + "\n")


def _write_json(path, filters):
    with path.open("w", encoding="utf-8") as handle:
        handle.write("[")
        first = True
        for row in _iter_rows(filters):
            if not first:
                handle.write(",")
            handle.write(json.dumps(dict(zip(HEADERS, row)), ensure_ascii=False))
            first = False
        handle.write("]")


def _write_txt(path, filters):
    with path.open("w", encoding="utf-8") as handle:
        handle.write("\t".join(HEADERS) + "\n")
        for row in _iter_rows(filters):
            handle.write("\t".join(str(value).replace("\t", " ") for value in row) + "\n")


def _write_xlsx(path, filters):
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("data")
    ws.append(HEADERS)
    for row in _iter_rows(filters):
        ws.append(row)
    wb.save(path)


def _write_docx(path, filters):
    from docx import Document

    document = Document()
    document.add_heading("Cofinets Data Intelligence Export", level=1)
    table = document.add_table(rows=1, cols=len(HEADERS))
    for cell, value in zip(table.rows[0].cells, HEADERS):
        cell.text = value
    for row in _iter_rows(filters):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    document.save(path)


def _find_pdf_font():
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf",
    ]
    return next((item for item in candidates if os.path.exists(item)), None)


def _write_pdf(path, filters):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4

    font_name = "Helvetica"
    font_path = _find_pdf_font()
    if font_path:
        font_name = "CDIUnicode"
        pdfmetrics.registerFont(TTFont(font_name, font_path))

    page_width, page_height = A4
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle("Cofinets Data Intelligence Export")
    y = page_height - 36
    pdf.setFont(font_name, 8)

    for row in _iter_rows(filters):
        text = " | ".join(f"{key}={value}" for key, value in zip(HEADERS, row))
        # Keep long JSON fields usable without making a single unbreakable line.
        for start in range(0, len(text), 115):
            pdf.drawString(24, y, text[start:start + 115])
            y -= 10
            if y < 30:
                pdf.showPage()
                pdf.setFont(font_name, 8)
                y = page_height - 36
    pdf.save()


WRITERS = {
    ExportJob.Format.CSV: _write_csv,
    ExportJob.Format.XLSX: _write_xlsx,
    ExportJob.Format.JSON: _write_json,
    ExportJob.Format.JSONL: _write_jsonl,
    ExportJob.Format.TXT: _write_txt,
    ExportJob.Format.DOCX: _write_docx,
    ExportJob.Format.PDF: _write_pdf,
}


def build_export(job):
    root = Path(getattr(settings, "EXPORT_ROOT", "/app/exports"))
    root.mkdir(parents=True, exist_ok=True)
    path = root / f"cdi-export-{job.pk}.{EXTENSIONS[job.format]}"
    writer = WRITERS.get(job.format)
    if writer is None:
        raise ValueError(f"unsupported export format: {job.format}")
    writer(path, job.filters)
    return path
